# -*- coding: utf-8 -*-
"""
Blankaga ruscha tarjima qo'shish (ikki tilli blanka) — IXTIYORIY.

Blanka Sozlamalari → "🌐 Blanka tili" → "Blankada ruscha tarjima ham chiqsin"
belgisi qo'yilgandagina ishlaydi (blanka_config.json: "blanka_ruscha": true).
Belgi o'chiq bo'lsa bu modul umuman chaqirilmaydi — blanka avvalgidek faqat
o'zbekcha chiqadi.

QANDAY ISHLAYDI:
  Blanka (python-docx Document) to'liq yasalib bo'lgach, SAQLASHDAN OLDIN
  ruscha_qoshish(doc) chaqiriladi. U hujjatdagi har bir jadval katagi va
  paragrafni ko'rib chiqadi: matni lug'atdagi o'zbekcha yozuvga AYNAN mos
  kelsa — ostiga kichik kulrang ruscha qator qo'shadi (vrach/UZI blankasidagi
  ".ru" uslubi: "Natija" → "Natija / Результат" ko'rinishida ikki qator).

  • Blanka yasovchi funksiyalarga tegilmaydi — shuning uchun o'zbekcha blanka
    buzilmaydi, ruscha faqat qo'shimcha qator sifatida qo'shiladi.
  • Lug'atda yo'q matn (bemor ismi, raqamlar, norma oraliqlari) o'zgarmaydi.
  • Matnida allaqachon kirill harfi bor katak (masalan lipid spektridagi
    "(Общий холестерин)") qayta tarjima qilinmaydi.

YANGI TAHLIL QO'SHILSA: uning blankadagi nomini (katta-kichik harf, apostrof
turi, ortiqcha bo'shliqlar ahamiyatsiz) quyidagi TARJIMA lug'atiga qo'shing.
"ИФА" yorlig'i avtomatik e'tiborsiz qoldiriladi.
"""
import re

try:
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
except Exception:  # python-docx yo'q bo'lsa — modul jim o'tadi
    Pt = RGBColor = qn = None


# ─────────────────────────────────────────────────────────────────────────────
# LUG'AT: o'zbekcha (blankada qanday chiqsa) → ruscha
# ─────────────────────────────────────────────────────────────────────────────
TARJIMA = {
    # ── Jadval sarlavhalari ────────────────────────────────────────────────
    "Tahlil nomi": "Наименование анализа",
    "Natija": "Результат",
    "Norma": "Норма",
    "Me'yor": "Норма",
    "O'lchov birligi": "Ед. измерения",
    "Birlik": "Ед. изм.",
    "Ko'rsatkich": "Показатель",
    "Ko'rsatgich": "Показатель",
    "Kod": "Код",
    "Qisqa nom": "Сокр.",
    "Daraja": "Степень",
    "Titr": "Титр",
    "Belgi (ramziy)": "Обозначение",
    "So'z bilan": "Словами",
    "Shaklli elementlar": "Форменные элементы",
    "Norma (Ayol)": "Норма (жен.)",
    "Norma (Erkak)": "Норма (муж.)",
    "Namuna olingan vaqt": "Время сбора",
    "Siydik miqdori (ml)": "Количество мочи (мл)",
    "Solishtirma og'irligi": "Удельный вес",

    # ── Bemor ma'lumotlari / poy ───────────────────────────────────────────
    "F.I.SH.": "Ф.И.О.",
    "F.I.SH": "Ф.И.О.",
    "Tug'ilgan sana": "Дата рождения",
    "Jins": "Пол",
    "Ayol": "Женский",
    "Erkak": "Мужской",
    "Laborant vrach": "Врач-лаборант",
    "Buyurtma №": "Заказ №",

    # ── Bo'lim (guruh) sarlavhalari ────────────────────────────────────────
    "QONNING BIOKIMYOVIY TAHLILI": "БИОХИМИЧЕСКИЙ АНАЛИЗ КРОВИ",
    "UMUMIY QON TAHLILI": "ОБЩИЙ АНАЛИЗ КРОВИ",
    "UMUMIY SIYDIK TAHLILI": "ОБЩИЙ АНАЛИЗ МОЧИ",
    "IMMUNOFERMENT ANALIZLAR": "ИММУНОФЕРМЕНТНЫЕ АНАЛИЗЫ (ИФА)",
    "EKSPRESS TESTLAR NATIJALARI": "РЕЗУЛЬТАТЫ ЭКСПРЕСС-ТЕСТОВ",
    "GEMOTOLOGIK TEKSHIRUVLAR": "ГЕМАТОЛОГИЧЕСКИЕ ИССЛЕДОВАНИЯ",
    "BOSHQA TESTLAR": "ДРУГИЕ ТЕСТЫ",
    "KOAGULOGRAMMA": "КОАГУЛОГРАММА",
    "QON GURUHI VA REZUS FAKTOR": "ГРУППА КРОВИ И РЕЗУС-ФАКТОР",
    "FIZIK-KIMYOVIY XOSSALARI": "ФИЗИКО-ХИМИЧЕСКИЕ СВОЙСТВА",
    "CHO'KMA MIKROSKOPIYASI": "МИКРОСКОПИЯ ОСАДКА",
    "BRUSELLYOZNI ANIQLASH UCHUN": "ДИАГНОСТИКА БРУЦЕЛЛЁЗА",
    "Reaksiya Xeddelsona va Rayta natijalari": "Результаты реакций Хеддельсона и Райта",
    "GINEKOLOGIK SURTMA": "ГИНЕКОЛОГИЧЕСКИЙ МАЗОК",
    "UROLOGIK SURTMA (MUJSKOY MAZOK)": "УРОЛОГИЧЕСКИЙ МАЗОК (МУЖСКОЙ)",
    "SPERMOGRAMMA": "СПЕРМОГРАММА",
    "NAJAS TAHLILI": "КОПРОГРАММА",
    "MIKROSKOPIYA NATIJALARI": "РЕЗУЛЬТАТЫ МИКРОСКОПИИ",
    "SIYDIKNING NECHIPORENKO TAHLILI": "АНАЛИЗ МОЧИ ПО НЕЧИПОРЕНКО",
    "SIYDIKNING ZIMNITSKIY TAHLILI": "АНАЛИЗ МОЧИ ПО ЗИМНИЦКОМУ",
    "QONDAN MAZOK (MORFOLOGIYA) NATIJALARI": "МАЗОК КРОВИ (МОРФОЛОГИЯ)",
    "TORCH INFEKSIYASIGA IgM VA IgG": "TORCH-ИНФЕКЦИИ IgM И IgG",
    "LIPID SPEKTRI": "ЛИПИДНЫЙ СПЕКТР",
    "BUYRAK PANELI": "ПОЧЕЧНАЯ ПАНЕЛЬ",
    "PEPSINOGEN": "ПЕПСИНОГЕН",
    "GLYUKOZAGA TOLERANTLIK TESTI (75 g glyukoza bilan)":
        "ГЛЮКОЗОТОЛЕРАНТНЫЙ ТЕСТ (с 75 г глюкозы)",
    "Leykogramma": "Лейкограмма",
    "Eritrotsitlar morfologiyasi": "Морфология эритроцитов",
    "Leykotsitlar morfologiyasi": "Морфология лейкоцитов",
    "Laborator xulosa": "Лабораторное заключение",
    "Periferik qon surtmasi tahlili": "Анализ мазка периферической крови",

    # ── Gematologiya ──────────────────────────────────────────────────────
    "Qonning umumiy tahlili": "Общий анализ крови",
    "Qondan mazok (morfologiya) ko'rish": "Мазок крови (морфология)",
    "Eritrotsitlar cho'kish tezligi (ECHT)": "Скорость оседания эритроцитов (СОЭ)",
    "Qonning ivish vaqti QIB": "Время свёртывания крови",
    "Qon guruhi va rezus faktorini aniqlash": "Определение группы крови и резус-фактора",
    "Qon guruhi": "Группа крови",
    "Rezus omil": "Резус-фактор",
    "Rezus musbat": "Резус положительный",
    "Rezus manfiy": "Резус отрицательный",
    "Birinchi": "Первая (0)",
    "Ikkinchi": "Вторая (A)",
    "Uchinchi": "Третья (B)",
    "To'rtinchi": "Четвёртая (AB)",
    "Leykotsitlar soni": "Количество лейкоцитов",
    "Limfotsitlar soni": "Количество лимфоцитов",
    "Monotsit + Eozinofil + Bazofil soni": "Моноциты + эозинофилы + базофилы (абс.)",
    "Neytrofillar soni": "Количество нейтрофилов",
    "Limfotsitlar ulushi": "Лимфоциты, %",
    "Monotsit + Eozinofil + Bazofil ulushi": "Моноциты + эозинофилы + базофилы, %",
    "Neytrofillar ulushi": "Нейтрофилы, %",
    "Eritrotsitlar soni": "Количество эритроцитов",
    "Gemoglobin": "Гемоглобин",
    "Gematokrit": "Гематокрит",
    "Eritrotsitlarning o'rtacha hajmi": "Средний объём эритроцита",
    "Eritrotsitdagi o'rtacha gemoglobin": "Среднее содержание гемоглобина в эритроците",
    "Eritrotsitdagi gemoglobin konsentratsiyasi": "Средняя концентрация гемоглобина в эритроците",
    "Eritrotsitlar hajmi variatsiyasi": "Ширина распределения эритроцитов (CV)",
    "Eritrotsitlar hajmi tarqalish kengligi": "Ширина распределения эритроцитов (SD)",
    "Trombotsitlar soni": "Количество тромбоцитов",
    "Trombotsitlarning o'rtacha hajmi": "Средний объём тромбоцита",
    "Trombotsitlar hajmi variatsiyasi": "Ширина распределения тромбоцитов",
    "Trombokrit": "Тромбокрит",
    "Neytrofil / Limfotsit nisbati": "Соотношение нейтрофилы / лимфоциты",
    "Trombotsit / Limfotsit nisbati": "Соотношение тромбоциты / лимфоциты",
    "Leykotsitlar (WBC)": "Лейкоциты (WBC)",
    "Mielositlar": "Миелоциты",
    "Metamielositlar": "Метамиелоциты",
    "Tayoqcha yadroli neytrofil": "Палочкоядерные нейтрофилы",
    "Segment yadroli neytrofil": "Сегментоядерные нейтрофилы",
    "Eozinofilllar": "Эозинофилы",
    "Eozinofillar": "Эозинофилы",
    "Bazofilllar": "Базофилы",
    "Bazofillar": "Базофилы",
    "Monotsitlar": "Моноциты",
    "Limfotsitlar": "Лимфоциты",
    "Plazmatik hujayralar": "Плазматические клетки",
    "Anizositoz": "Анизоцитоз",
    "Poikilositoz": "Пойкилоцитоз",
    "Gipoxromiya": "Гипохромия",
    "Mikrotsitoz": "Микроцитоз",
    "Makrotsitoz": "Макроцитоз",

    # ── Bioximiya ─────────────────────────────────────────────────────────
    "Albumin": "Альбумин",
    "Umumiy oqsil": "Общий белок",
    "Glyukoza": "Глюкоза",
    "Glyukozaga tolerantlik testi": "Глюкозотолерантный тест",
    "Glyukoza-natoshchak": "Глюкоза натощак",
    "Glyukoza 1 soat o'tgach": "Глюкоза через 1 час",
    "Glyukoza 2 soat o'tgach": "Глюкоза через 2 часа",
    "Alanin aminotransferaza (ALT)": "Аланинаминотрансфераза (АЛТ)",
    "Aspartat aminotransferaza (AST)": "Аспартатаминотрансфераза (АСТ)",
    "Kreatinin": "Креатинин",
    "Mochevina": "Мочевина",
    "Umumiy xolesterin (TC)": "Общий холестерин",
    "Trigliserid-TG": "Триглицериды",
    "Laktatdegidrogrenaza (LDG)": "Лактатдегидрогеназа (ЛДГ)",
    "Gamma-glutamil transferaza (GGT)": "Гамма-глутамилтрансфераза (ГГТ)",
    "Kalsiy Ca": "Кальций",
    "Kalsiy": "Кальций",
    "Kaliy K": "Калий",
    "Temir Fe": "Железо",
    "Siydik kislota (Mochevaya kislota)": "Мочевая кислота",
    "Magniy Mg": "Магний",
    "Natriy Na": "Натрий",
    "Ishqoriy fosfataza (IF)": "Щелочная фосфатаза (ЩФ)",
    "Timol proba": "Тимоловая проба",
    "Bilirubin": "Билирубин",
    "BILIRUBIN (umumiy, bog'langan, erkin)": "БИЛИРУБИН (общий, прямой, непрямой)",
    "Umumiy Bilirubin": "Билирубин общий",
    "Bog'langan Bilirubin": "Билирубин прямой",
    "Erkin Bilirubin": "Билирубин непрямой",
    "Alfa-amilaza": "Альфа-амилаза",
    "Alfa-amilaza pankreaticheskiy": "Альфа-амилаза панкреатическая",
    "REVMOPROBA to'liq (avtomat)": "Ревмопробы (полный)",
    "REVMOPROBA to'liq (qo'lda)": "Ревмопробы (полный)",
    "RF (Revmatoidniy faktor) avtomat": "Ревматоидный фактор (РФ)",
    "CRP (S-reaktivniy belok) avtomat": "С-реактивный белок (СРБ)",
    "ASLO (Antistreptolizin-O) avtomat": "Антистрептолизин-О (АСЛО)",
    "RF (Revmatoidniy faktor)": "Ревматоидный фактор (РФ)",
    "CRP (S-reaktivniy belok)": "С-реактивный белок (СРБ)",
    "ASLO (Antistreptolizin-O)": "Антистрептолизин-О (АСЛО)",
    "LDL-Past zichlikli lipoprotein": "Липопротеины низкой плотности (ЛПНП)",
    "HDL-Yuqori zichlikli lipoprotein": "Липопротеины высокой плотности (ЛПВП)",
    "Xolinesteraza (CHE)": "Холинэстераза",
    "Mikroalbumin": "Микроальбумин",
    "Mikroalbumin/Kreatinin nisbati": "Соотношение микроальбумин/креатинин",

    # ── Koagulogramma ─────────────────────────────────────────────────────
    "KOAGULOGRAMMA to'liq": "Коагулограмма (полная)",
    "Protrombin Time (PTI, PT/MNO)": "Протромбиновое время (ПТИ, МНО)",
    "PT sekundalarda": "ПВ в секундах",
    "PT bo'yicha Kviku (%)": "ПТИ по Квику (%)",
    "MNO / INR": "МНО",
    "Fibrinogen": "Фибриноген",
    "ACHTV": "АЧТВ",
    "Trombinovoe vremya (TT)": "Тромбиновое время (ТВ)",
    "Trombin vaqti (TT)": "Тромбиновое время (ТВ)",

    # ── Siydik ────────────────────────────────────────────────────────────
    "Siydikning umumiy tahlili": "Общий анализ мочи",
    "Siydikning Nicheporenko tahlili": "Анализ мочи по Нечипоренко",
    "Siydikning Zimnitskiy tahlili": "Анализ мочи по Зимницкому",
    "Miqdori": "Количество",
    "Rangi": "Цвет",
    "Tiniqligi": "Прозрачность",
    "Leykotsit": "Лейкоциты",
    "Leykotsitlar": "Лейкоциты",
    "Keton": "Кетоны",
    "Nitritlar": "Нитриты",
    "Urobilinogen": "Уробилиноген",
    "Oqsil": "Белок",
    "Nisbiy zich.": "Относ. плотность",
    "Qon": "Кровь",
    "Askorbin": "Аскорбиновая кислота",
    "Epiteliy: yassi": "Эпителий плоский",
    "Epiteliy: o'tuvchi": "Эпителий переходный",
    "Epiteliy: buyrak": "Эпителий почечный",
    "Eritrotsitlar: o'zgargan": "Эритроциты изменённые",
    "Eritrotsitlar: o'zgarmagan": "Эритроциты неизменённые",
    "Eritrotsitlar": "Эритроциты",
    "Silindrlar": "Цилиндры",
    "Silindirlar": "Цилиндры",
    "Shilliq": "Слизь",
    "Tuzlar": "Соли",
    "Bakteriyalar": "Бактерии",
    "Zamburug'lar": "Грибы",
    "Trixomonada": "Трихомонады",
    "Kunduzgi diurez": "Дневной диурез",
    "Tungi diurez": "Ночной диурез",
    "Jami sutkalik diurez": "Суточный диурез",

    # ── Ekspress / boshqa testlar ─────────────────────────────────────────
    "Troponin T": "Тропонин T",
    "Troponin I": "Тропонин I",
    "Troponin I test": "Тропонин I (экспресс)",
    "Troponin I miqdoriy": "Тропонин I (количественный)",
    "Axlatdan yashirin qon testi": "Скрытая кровь в кале",
    "Zamburug' tekshiruvi": "Исследование на грибы",
    "Qo'tir kanasi aniqlash": "Исследование на чесоточного клеща",
    "Demodikozni aniqlash": "Исследование на демодекоз",
    "Gepatit B (HBsAg, ekspress)": "Гепатит B (HBsAg, экспресс)",
    "Gepatit C (HCV-Ab, ekspress)": "Гепатит C (анти-HCV, экспресс)",
    "Brutselyoz (Reaktsiya Xedelson + Rayta)": "Бруцеллёз (реакции Хеддельсона и Райта)",
    "Reaksiya Xeddelsona": "Реакция Хеддельсона",
    "Reaksiya Rayta": "Реакция Райта",
    "Helicobakter pylori AB (ekspress)": "Helicobacter pylori (антитела, экспресс)",
    "H.Pylori Test (axlatdan)": "H. pylori (антиген в кале)",
    "TORCH test IgM va IgG (to'liq)": "TORCH-инфекции IgM и IgG",
    "Toxoplasma (TOKSOPLAZMOZ) Ekspress test": "Токсоплазмоз (экспресс-тест)",
    "Rubella (QIZILCHA) Ekspress test": "Краснуха (экспресс-тест)",
    "Cytomegalovirus (SITOMYEGALOVIRUS) Ekspress test": "Цитомегаловирус (экспресс-тест)",
    "Herpes 1 (ODDIY GERPYES 1 tipi) Ekspress test": "Герпес 1 типа (экспресс-тест)",
    "Herpes 2 (ODDIY GERPYES 2 tipi) Ekspress test": "Герпес 2 типа (экспресс-тест)",
    "RW sifilis (ekspress)": "Сифилис RW (экспресс)",
    "OIV": "ВИЧ",
    "Ginekologik surtma": "Гинекологический мазок",
    "Spermogramma": "Спермограмма",
    "Mujskoy mazok (Erkaklar surtmasi)": "Мужской мазок",
    "Najas tahlili (Gijja uchun)": "Анализ кала (на гельминты)",

    # ── Ginekologik / urologik surtma ─────────────────────────────────────
    "Qin (V)": "Влагалище (V)",
    "Bachadon bo'yni (C)": "Шейка матки (C)",
    "Siydik yo'li (U)": "Уретра (U)",
    "Yassi epiteliy": "Плоский эпителий",
    "Epiteliy": "Эпителий",
    "Gonokokklar": "Гонококки",
    "Atipik (klyucheviy) xujayralar": "Ключевые клетки",
    "Laktobatsillalar (tayoqcha Doderleyna, gramm musbat tayoqcha)": "Лактобациллы (палочки Додерлейна)",
    "Candida (Aitqi zamburug'i)": "Candida (дрожжевые грибы)",
    "Anaerob bakteriyalar (gramm manfiy tayoqcha), kokklar": "Анаэробные бактерии, кокки",
    "Kokklar mikroflora": "Кокковая микрофлора",
    "Hujayra ichi gram manfiy diplakokklari (Gonokokklar)": "Внутриклеточные грамотрицательные диплококки (гонококки)",
    "Hujayradan tashqari gramm manfiy palochka va kokklar": "Внеклеточные грамотрицательные палочки и кокки",
    "Hujayradan tashqari gramm musbat kokklar": "Внеклеточные грамположительные кокки",
    "Xlamidiya": "Хламидии",
    "Ureaplazma": "Уреаплазма",

    # ── Spermogramma ──────────────────────────────────────────────────────
    "Saqlanish vaqti": "Период воздержания",
    "Xajmi (ml)": "Объём (мл)",
    "Xidi": "Запах",
    "Suyulish vaqti (daq.)": "Время разжижения (мин)",
    "Ilashuvchanglik (mm)": "Вязкость (мм)",
    "Spermatozoidlar a) 1 mlda": "Сперматозоиды в 1 мл",
    "Spermatozoidlar b) umumiy xajmida": "Сперматозоиды в эякуляте",
    "Harakati: a) faol (%)": "Подвижные активные (%)",
    "Harakati: b) sust (%)": "Малоподвижные (%)",
    "Harakati: v) harakatsiz (%)": "Неподвижные (%)",
    "Tirik spermatozoidlar (%)": "Живые сперматозоиды (%)",
    "Patologik shakllar (%)": "Патологические формы (%)",
    "Spermatogen epiteliy": "Сперматогенный эпителий",
    "Letsitin donachalar": "Лецитиновые зёрна",
    "Agglyutinatsiya": "Агглютинация",
    "Fruktoza (mkmol/eyak.)": "Фруктоза (мкмоль/эяк.)",
    "Limon kislotasi (mkmol/eyak.)": "Лимонная кислота (мкмоль/эяк.)",

    # ── Najas (koprogramma) ───────────────────────────────────────────────
    "Sutkalik miqdori": "Суточное количество",
    "Namuna miqdori": "Количество образца",
    "Konsistensiyasi": "Консистенция",
    "Shakli": "Форма",
    "Reaksiyasi (pH)": "Реакция (pH)",
    "O'simlik kletchatkasi": "Растительная клетчатка",
    "Kraxmal": "Крахмал",
    "Yodofil flora": "Йодофильная флора",
    "Shilliq, epiteliy": "Слизь, эпителий",
    "Neytral yog'": "Нейтральный жир",
    "Shilliq (mikroskopiya)": "Слизь (микроскопия)",
    "Sodda xayvonlar (Protozoa)": "Простейшие",
    "Gijja tuxumlari": "Яйца гельминтов",

    # ── IFA ───────────────────────────────────────────────────────────────
    "Pepsinogen I (PGI) IFA": "Пепсиноген I",
    "Pepsinogen II (PGII) IFA": "Пепсиноген II",
    "PG I / PG II Nisbati": "Соотношение PG I / PG II",
    "CA 72-4 (oshqozon onkomarkeri)": "CA 72-4 (онкомаркер желудка)",
    "CA15-3 (ko'krak saratoni belgisi)": "CA 15-3 (онкомаркер молочной железы)",
    "Gepatit B (HBsAg, IFA)": "Гепатит B (HBsAg, ИФА)",
    "Gepatit C (Anti-HCV, IFA)": "Гепатит C (анти-HCV, ИФА)",
    "Gepatit D (D-antitela, IFA)": "Гепатит D (антитела, ИФА)",
    "Gepatit A (IgG, IFA)": "Гепатит A (IgG, ИФА)",
    "Helicobakter pylori IgG IFA": "Helicobacter pylori IgG (ИФА)",
    "IgE umumiy (Allergiya uchun)": "IgE общий",
    "Askarida IgG IFA": "Аскариды IgG (ИФА)",
    "Lyambliya-antitela IFA": "Лямблии, антитела (ИФА)",
    "Vitamin D (25-OH)": "Витамин D (25-OH)",
    "Vitamin B12": "Витамин B12",
    "Ferritin": "Ферритин",
    "TTG (TSH) IFA": "ТТГ (тиреотропный гормон)",
    "T4 erkin IFA": "Т4 свободный",
    "T3 erkin IFA": "Т3 свободный",
    "AT-TPO IFA": "Антитела к ТПО",
    "LG (LH) IFA": "ЛГ (лютеинизирующий гормон)",
    "FSG (FSH) IFA": "ФСГ (фолликулостимулирующий гормон)",
    "Prolaktin IFA": "Пролактин",
    "Estradiol IFA": "Эстрадиол",
    "Progesteron IFA": "Прогестерон",
    "Testesteron IFA": "Тестостерон",
    "Testesteron Erkin IFA": "Тестостерон свободный",
    "Umumiy PSA IFA": "ПСА общий",
    "Insulin IFA": "Инсулин",
    "Kortizol IFA": "Кортизол",
    "XGCh (Homiladorlik gormoni) IFA": "ХГЧ (хорионический гонадотропин)",
    "Anti HBsAg (Antitelo) immunitet IFA": "Анти-HBs (антитела, иммунитет)",
    "Troponin I miqdoriy IFA": "Тропонин I (количественный)",

    # ── Tez-tez uchraydigan natija so'zlari ───────────────────────────────
    "Manfiy (-)": "Отрицательно (-)",
    "Manfiy": "Отрицательно",
    "(0) manfiy": "(0) отрицательно",
    "Musbat (+)": "Положительно (+)",
    "Musbat (++)": "Положительно (++)",
    "Musbat (+++)": "Положительно (+++)",
    "Musbat (+/-)": "Слабоположительно (+/-)",
    "Musbat": "Положительно",
    "Aniqlanmadi": "Не обнаружено",
    "Aniqlandi": "Обнаружено",
    "Normal": "Норма",
    "Tiniq": "Прозрачная",
    "Xira": "Мутная",
    "Biroz xira": "Слегка мутная",
    "Sariq": "Жёлтый",
    "Och sariq": "Светло-жёлтый",
    "Qo'ng'ir sut": "Коричневый",
    "Jigarrang": "Коричневый",
    "Xom kashtan": "Бурый",
    "Qon aralash": "С примесью крови",
    "Neytral": "Нейтральная",
    "Ko'p miqdorda": "В большом количестве",
    "Yumshoq": "Мягкий",
    "Suyuq": "Жидкий",
    "Silindrsimon": "Цилиндрическая",
    "Shakllangan (yumshoq)": "Оформленный (мягкий)",
    "Shakllangan (silindrsimon)": "Оформленный (цилиндрический)",
    "Zamburug' aniqlandi": "Грибы обнаружены",
    "Zamburug' aniqlanmadi": "Грибы не обнаружены",
    "Demodikoz aniqlanmadi": "Демодекс не обнаружен",
    "Patologik o'zgarishlar aniqlanmadi.": "Патологических изменений не выявлено.",
    "Bor": "Есть",
    "ko'ruv maydonida": "в поле зрения",
    "preparatda": "в препарате",

    # ── Buyrak paneli (hisoblangan ko'rsatkichlar) ────────────────────────
    "Mochevina azoti (BUN) (hisoblangan)": "Азот мочевины (BUN), расчётный",
    "Kreatinin / Mochevina nisbati (hisoblangan)": "Соотношение креатинин / мочевина, расчётное",
    "Umumiy azot miqdori (hisoblangan)": "Общий азот, расчётный",
    "GFR — Klubochkaviy filtrlash tezligi (CKD-EPI formulasi, avtomatik hisoblangan)":
        "СКФ — скорость клубочковой фильтрации (CKD-EPI, расчётная)",
    "Yo'q": "Нет",
}

# Paragraf SHU so'z bilan boshlansa — faqat boshlanishining tarjimasi qo'shiladi
# (davomida sana/raqam turadi: "Chop etilgan vaqt: 24.09.2026 08:54").
PREFIKS = {
    "Chop etilgan vaqt": "Время печати",
    "Avtomatik gematologik analizator": "Автоматический гематологический анализатор",
    "Kimyoviy tahlil (URIT-50 Analizator)": "Химический анализ мочи (анализатор URIT-50)",
    "Menstrual siklning": "День менструального цикла",
}

# NORMA QATORLARI ICHIDAGI so'zlar (jins / yosh / faza / trimestr / natija so'zi).
# Norma qatori "Erkak: 44.0-115.0" ko'rinishida bo'ladi — butun qatorni
# tarjima qilib bo'lmaydi (raqamlar o'zgaradi), shuning uchun so'z topilgan
# joyning O'ZIGA kichik kulrang ruscha qo'shiladi:
#   "Erkak: 44.0-115.0"          → "Erkak / муж.: 44.0-115.0"
#   "12-19 yosh Erkak: 1.0-38.0" → "12-19 yosh Erkak / лет, муж.: 1.0-38.0"
# Yonma-yon (faqat bo'shliq bilan ajralgan) so'zlar bitta izohga birlashadi.
# TARTIB MUHIM: uzunroq ibora oldin turishi kerak ("Emizikli yosh" > "yosh").
# Qolipdagi "'" har qanday apostrof turiga mos keladi; (\d...) guruhi {1} ga tushadi.
NORMA_SOZLAR = [
    # Yosh oraliqlari (raqam bilan birga)
    (r"(\d+)\s+oydan\s+(\d+)\s+yoshgacha", "от {1} мес. до {2} лет"),
    (r"(\d+)\s+yoshgacha(?:\s+yosh)?", "до {1} лет"),
    (r"(\d+)\s*-\s*trimestr", "{1}-й триместр"),
    (r"(\d+)\s+trimestr", "{1}-й триместр"),
    (r"(\d+)\s*-\s*Hafta", "{1}-я нед."),
    (r"(\d+(?:\s*-\s*\d+)?)\s+Hafta", "{1} нед."),
    (r"(\d+[ab]?)\s*-\s*bosqich", "стадия {1}"),
    (r"(\d[\d.,]*)\s*gacha", "до {1}"),
    (r"(\d+)\s*soatdan\s+keyin\s+takrorlash", "повторить через {1} ч"),
    # Toifalar
    ("Emizikli yosh", "лактация"),
    ("Emizikli", "лактация"),
    ("Yangi tug'ilgan", "новорождённые"),
    ("Chaqaloq", "новорождённые"),
    ("Erkaklarda", "у мужчин"),
    ("Ayollarda", "у женщин"),
    ("Erkak", "муж."),
    ("Ayol", "жен."),
    ("Follikulyar", "фолликулярная фаза"),
    ("Ovulyatsiya", "овуляция"),
    ("Lyutein", "лютеиновая фаза"),
    ("Menopauza", "менопауза"),
    ("qandli diabet", "сахарный диабет"),
    ("yosh", "лет"),
    ("oy", "мес."),
    # Vitamin D bosqichlari
    ("Juda kuchli yetmaslik", "выраженный дефицит"),
    ("Og'ir yetishmovchilik", "тяжёлый дефицит"),
    ("Yengil yetishmovchilik", "лёгкий дефицит"),
    ("Norma (suboptimal)", "норма (субоптимальная)"),
    ("Norma (optimal)", "норма (оптимальная)"),
    ("Yuqori norma", "высокая норма"),
    ("Yuqori toksik emas", "высокий, не токсичный"),
    ("Potensial toksiklik", "потенциально токсичный"),
    # GFR bosqichlari
    ("Buyrak yetishmovchiligi", "почечная недостаточность"),
    ("Yengil pasayish", "лёгкое снижение"),
    ("O'rtacha pasayish", "умеренное снижение"),
    ("O'rtacha[–-]og'ir", "умеренно-тяжёлое"),
    ("Og'ir", "тяжёлое"),
    # Natija so'zlari (norma matni ichida)
    (r"Manfiy(?:\s*\(\s*-\s*\))?", "отриц."),
    (r"Musbat(?:\s*\(\s*\+[+/-]*\s*\))?", "полож."),
    ("Shubhali (chegaraviy)", "сомнительно (пограничное)"),
    ("Shubhali", "сомнительно"),
    ("qayta tekshirish", "повторить исследование"),
    ("Aniqlanmadi", "не обнаружено"),
    ("Aniqlandi", "обнаружено"),
    ("Normal", "норма"),
    ("Norma", "норма"),
    # Ivish vaqti, spermogramma
    ("Boshlandi", "начало"),
    ("Tugadi", "окончание"),
    ("sutka", "сут."),
    (r"daq\.", "мин"),
    # Kolontitul (Sana: ...  Vaqt: ...   1 / 3 sahifa)
    (r"Sana(?=\s*:)", "Дата"),
    (r"Vaqt(?=\s*:)", "Время"),
    ("sahifa", "(стр.)"),   # "(" bilan boshlansa "/" qo'yilmaydi
]

_APOS_CLS = "['‘’ʻʼ`´]"


def _norma_qolip(q):
    # Oddiy ibora (regex belgisisiz) → escape; apostrof → istalgan apostrof turi
    if not re.search(r"[\\(\[?]", q) or q.startswith(("Norma (", "Shubhali (")):
        q = re.escape(q)
    q = q.replace("\\'", "'").replace("'", _APOS_CLS)
    return re.compile(r"(?<![A-Za-z‘’ʻʼ'])" + q + r"(?![A-Za-z])",
                      re.IGNORECASE)


_NORMA_QOLIPLAR = [(_norma_qolip(q), ru) for q, ru in NORMA_SOZLAR]

_KIRILL_RE = re.compile(r"[А-Яа-яЁё]")
_APOS_RE = re.compile(r"[‘’ʻʼ`´]")


def _norm(s):
    """Taqqoslash uchun: apostrof turlari, 'ИФА' yorlig'i, ortiqcha bo'shliq,
    oxiridagi ':' va katta-kichik harf ahamiyatsiz."""
    s = _APOS_RE.sub("'", s or "")
    s = s.replace("ИФА", " ")
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"\(\s+", "(", re.sub(r"\s+\)", ")", s))
    s = s.strip().rstrip(":").strip()
    return s.lower()


_TARJIMA_N = {_norm(k): v for k, v in TARJIMA.items()}
_PREFIKS_N = sorted(((_norm(k), v) for k, v in PREFIKS.items()),
                    key=lambda kv: -len(kv[0]))


def _kirill_bor(s):
    """'ИФА' yorlig'idan tashqari kirill harfi bormi (allaqachon ikki tilli)."""
    return bool(_KIRILL_RE.search((s or "").replace("ИФА", "")))


def tarjima(matn):
    """O'zbekcha matnning ruscha tarjimasi yoki None."""
    if not matn or _kirill_bor(matn):
        return None
    n = _norm(matn)
    if not n:
        return None
    ru = _TARJIMA_N.get(n)
    if ru:
        return ru
    # "... IFA" / "... avtomat" qo'shimchasisiz ham qidiramiz
    for suf in (" ifa", " avtomat"):
        if n.endswith(suf):
            ru = _TARJIMA_N.get(n[: -len(suf)].strip())
            if ru:
                return ru
    return None


def _prefiks_tarjima(matn):
    if not matn or _kirill_bor(matn):
        return None
    n = _norm(matn)
    for k, v in _PREFIKS_N:
        if n.startswith(k) and n != k:
            return v
    return None


def _run_rgb(run):
    try:
        c = run.font.color
        if c is not None and c.type is not None and c.rgb is not None:
            return c.rgb
    except Exception:
        pass
    return None


def _qoshish(paragraph, ru):
    """Paragraf oxiriga: yangi qator + kichik ruscha matn."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    asos = runs[-1] if runs else None

    size = None
    bold = False
    rgb = None
    font_name = "Times New Roman"
    if asos is not None:
        try:
            if asos.font.size:
                size = asos.font.size.pt
        except Exception:
            pass
        bold = bool(asos.bold)
        rgb = _run_rgb(asos)
        if asos.font.name:
            font_name = asos.font.name
    if size is None:
        size = 11.0
    # Ruscha qator: asl matndan kichikroq (vrach blankasidagi 7.5pt ga yaqin)
    ru_size = max(7.0, min(round(size * 0.72 * 2) / 2, 10.5))

    br = paragraph.add_run()
    br.add_break()
    if asos is not None:
        try:
            br.font.size = asos.font.size
        except Exception:
            pass
    r = paragraph.add_run(ru)
    r.font.size = Pt(ru_size)
    r.font.name = font_name
    try:
        rpr = r._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), font_name)
    except Exception:
        pass
    # Sarlavhalar (qalin) — qalin qoladi; oddiy matn — kulrang
    r.bold = bold
    if rgb is not None and str(rgb).upper() not in ("000000", "auto"):
        r.font.color.rgb = rgb      # rangli sarlavha (oq/binafsha) — o'sha rang
    else:
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def _norma_topish(matn):
    """Norma so'zlarini topadi → [(boshi, oxiri, ruscha), ...] (bir-birini yopmaydi).
    Yonma-yon (faqat bo'shliq bilan ajralgan) topilmalar bitta izohga birlashadi."""
    topilgan = []
    band = [False] * len(matn)
    for rx, ru in _NORMA_QOLIPLAR:
        for m in rx.finditer(matn):
            a, b = m.span()
            if a == b or any(band[a:b]):
                continue
            for i in range(a, b):
                band[i] = True
            t = ru
            for gi, g in enumerate(m.groups(), 1):
                t = t.replace("{%d}" % gi, re.sub(r"\s+", "", g or ""))
            topilgan.append((a, b, t))
    topilgan.sort()
    guruhlar = []
    for a, b, t in topilgan:
        if guruhlar and not matn[guruhlar[-1][1]:a].strip():
            pa, pb, pt = guruhlar[-1]
            guruhlar[-1] = (pa, b, pt + ", " + t)
        else:
            guruhlar.append((a, b, t))
    return guruhlar


def _run_boshqa_nusxa(run_el, matn):
    """Run elementining (rPr formati bilan) nusxasi, boshqa matn bilan."""
    import copy
    yangi = copy.deepcopy(run_el)
    for ch in list(yangi):
        if ch.tag != qn("w:rPr"):
            yangi.remove(ch)
    from docx.text.run import Run
    Run(yangi, None).text = matn
    return yangi


def _norma_izoh_qoshish(paragraph):
    """Norma qatori ichidagi so'zlar yoniga " / ruscha" (kichik, kulrang) qo'shadi.
    Qalin (bemorga mos) qator — izoh ham qalin. True = biror narsa qo'shildi."""
    from docx.text.run import Run
    runs = paragraph.runs
    matn = "".join(r.text for r in runs)
    guruhlar = _norma_topish(matn)
    if not guruhlar:
        return False
    izohlar = set()
    # O'ngdan chapga — oldingi pozitsiyalar siljimasin
    for _a, pos, ru in reversed(guruhlar):
        runs = paragraph.runs
        start = 0
        for r in runs:
            rt = r.text
            if start < pos <= start + len(rt):
                k = pos - start
                if k < len(rt):  # run'ni ikkiga bo'lamiz: [bosh][izoh][dum]
                    dum = _run_boshqa_nusxa(r._r, rt[k:])
                    r.text = rt[:k]
                    r._r.addnext(dum)
                izoh_el = _run_boshqa_nusxa(r._r, (" " if ru.startswith("(") else " / ") + ru)
                r._r.addnext(izoh_el)
                izohlar.add(izoh_el)
                izoh = Run(izoh_el, paragraph)
                size = r.font.size.pt if r.font.size else 11.0
                izoh.font.size = Pt(max(7.0, min(round(size * 0.72 * 2) / 2, 10.5)))
                rgb = _run_rgb(r)
                if rgb is None or str(rgb).upper() == "000000":
                    izoh.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                break
            start += len(rt)
    _oraliq_uzilmas(paragraph, izohlar)
    return True


_ORALIQ_RE = re.compile(r"(?<=\d)-(?=[\d.])")


def _oraliq_uzilmas(paragraph, izohlar):
    """Norma oralig'idagi chiziqchani (0.05-0.7) Word'ning uzilmas chiziqchasiga
    (w:noBreakHyphen) almashtiradi — ruscha izoh qo'shilib qator uzayganda raqam
    "0.05-" / "0.7" bo'lib ikki qatorga bo'linib ketmasin. Ko'rinishi o'zgarmaydi."""
    for r in paragraph.runs:
        if r._r in izohlar:
            continue
        t = r.text
        if not _ORALIQ_RE.search(t) or "\n" in t or "\t" in t:
            continue
        el = r._r
        for ch in list(el):
            if ch.tag != qn("w:rPr"):
                el.remove(ch)
        qismlar = _ORALIQ_RE.split(t)
        for i, q in enumerate(qismlar):
            if i:
                el.append(el.makeelement(qn("w:noBreakHyphen"), {}))
            if q:
                te = el.makeelement(qn("w:t"), {})
                te.text = q
                te.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                el.append(te)


def _paragraflar(paragraphs, bajarilgan, kolontitul=False):
    for p in paragraphs:
        if p._p in bajarilgan:
            continue
        t = p.text
        if not t.strip():
            continue
        ru = tarjima(t) or _prefiks_tarjima(t)
        if ru:
            _qoshish(p, ru)
            bajarilgan.add(p._p)
        elif (kolontitul or re.search(r"\d", t)) and not _kirill_bor(t):
            # (Kolontitulda sahifa raqami Word maydoni — matnda raqam ko'rinmaydi)
            # Norma qatori ("Erkak: 44.0-115.0", "<3 yosh: ...", "1-trimestr: ...")
            # yoki kolontitul ("Sana: ... Vaqt: ...") — so'z yoniga ruscha izoh
            try:
                if _norma_izoh_qoshish(p):
                    bajarilgan.add(p._p)
            except Exception as e:
                print(f"[RUSCHA] norma izohi: {e}")


def _jadvallar(tables, bajarilgan, kolontitul=False):
    for tbl in tables:
        for row in tbl.rows:
            korilgan = set()
            for cell in row.cells:
                # Birlashtirilgan (merge) kataklar bir necha marta qaytadi.
                # DIQQAT: id() emas — lxml proxy obyektlari yo'qolgach id qayta
                # ishlatiladi va boshqa katak "ko'rilgan" deb o'tkazib yuboriladi.
                if cell._tc in korilgan:
                    continue
                korilgan.add(cell._tc)
                paras = [p for p in cell.paragraphs if p.text.strip()]
                if len(paras) > 1:
                    # Butun katak bitta nom bo'lsa (bir necha paragrafga bo'lingan)
                    ru = tarjima(cell.text)
                    if ru:
                        _qoshish(paras[-1], ru)
                        for p in paras:
                            bajarilgan.add(p._p)
                _paragraflar(cell.paragraphs, bajarilgan, kolontitul)
                if cell.tables:
                    _jadvallar(cell.tables, bajarilgan, kolontitul)


def ruscha_qoshish(doc):
    """Tayyor blankaga ruscha tarjima qatorlarini qo'shadi (joyida o'zgartiradi).
    Xatolik bo'lsa ham blankani buzmaydi — shunchaki shu joyni o'tkazib yuboradi."""
    if doc is None or Pt is None:
        return doc
    bajarilgan = set()
    try:
        _paragraflar(doc.paragraphs, bajarilgan)
        _jadvallar(doc.tables, bajarilgan)
    except Exception as e:
        print(f"[RUSCHA] asosiy qism: {e}")
    try:
        for sec in doc.sections:
            for hf in (sec.header, sec.footer):
                if hf is None or hf.is_linked_to_previous:
                    continue
                _paragraflar(hf.paragraphs, bajarilgan, kolontitul=True)
                _jadvallar(hf.tables, bajarilgan, kolontitul=True)
    except Exception as e:
        print(f"[RUSCHA] kolontitul: {e}")
    return doc
