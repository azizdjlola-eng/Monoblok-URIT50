"""Xodimlar qo'llanmalari: Markdown (.md) -> Word (.docx) + PDF.

Manba — shu papkadagi *.md fayllar (git'da saqlanadi, uyda ham ishda ham bir xil).
Natija — shu papkadagi .docx/.pdf va OneDrive\\AzizMedLine_Qollanmalar ga nusxa.

Ishlatish:
    python qollanma_yasash.py              # barcha .md
    python qollanma_yasash.py qon_olish_qollanma.md
PDF uchun kompyuterda MS Word bo'lishi kerak (bo'lmasa faqat .docx).

Qo'llab-quvvatlanadigan Markdown: # / ## / ### sarlavha, paragraf, "- " ro'yxat,
"1. " raqamli ro'yxat, "- [ ] " belgilash katakchasi, | jadval |, **qalin**, `kod`.
"""
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PAPKA = os.path.dirname(os.path.abspath(__file__))
ONEDRIVE = os.path.join(os.environ.get("OneDrive") or os.path.expanduser("~\\OneDrive"),
                        "AzizMedLine_Qollanmalar")
NAVY = RGBColor(0x1A, 0x3A, 0x5C)
NAVY_HEX = "1A3A5C"
ZEBRA = "EEF3F8"
SHRIFT = "Times New Roman"


def _shrift(style, size, bold=None, color=None):
    style.font.name = SHRIFT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), SHRIFT)


def _hujjat():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    _shrift(doc.styles["Normal"], 11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    for name, size in (("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)):
        h = doc.styles[name]
        _shrift(h, size, bold=True, color=NAVY)
        h.paragraph_format.space_before = Pt(0 if name == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    return doc


def _matn(par, text, size=None, color=None, bold_all=False):
    """**qalin** va `kod` -> qalin; \\_ -> _."""
    text = text.replace("\\_", "_")
    for p in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not p:
            continue
        bold = bold_all
        if p.startswith("`") or p.startswith("**"):
            p = p.strip("`").strip("*")
            bold = True
        r = par.add_run(p)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def _fon(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_)
    tcPr.append(s)


def _jadval(doc, qatorlar):
    def katak(q):
        return [c.strip() for c in q.strip().strip("|").split("|")]
    sarlavha = katak(qatorlar[0])
    tana = [katak(q) for q in qatorlar[2:]]
    n = len(sarlavha)
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(sarlavha):
        c = t.rows[0].cells[i]
        _fon(c, NAVY_HEX)
        _matn(c.paragraphs[0], h, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(th)   # sarlavha har sahifada takrorlansin
    for ri, q in enumerate(tana):
        cells = t.add_row().cells
        for i in range(n):
            _matn(cells[i].paragraphs[0], q[i] if i < len(q) else "", size=10, bold_all=(i == 0))
            if ri % 2 == 1:
                _fon(cells[i], ZEBRA)
    # ustun kengligi — matn uzunligiga qarab (sahifa eni 17.4 sm)
    uz = [max(len(sarlavha[i]), *(len(q[i]) if i < len(q) else 0 for q in tana)) for i in range(n)]
    uz = [min(max(u, 6), 60) for u in uz]
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i in range(n):
            row.cells[i].width = Cm(17.4 * uz[i] / sum(uz))
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def md_dan_docx(md_yol, docx_yol):
    doc = _hujjat()
    qatorlar = open(md_yol, encoding="utf-8").read().splitlines()
    i = 0
    birinchi_paragraf = True
    while i < len(qatorlar):
        q = qatorlar[i]
        s = q.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            blok = []
            while i < len(qatorlar) and qatorlar[i].strip().startswith("|"):
                blok.append(qatorlar[i])
                i += 1
            _jadval(doc, blok)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif s.startswith("- [ ]") or s.startswith("- [x]"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run("☑  " if s.startswith("- [x]") else "☐  ")
            r.font.size = Pt(13)
            _matn(p, s[5:].strip())
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _matn(p, s[2:])
        elif re.match(r"^\d+\.\s", s):
            # raqamni o'zimiz yozamiz — Word ro'yxatlari bo'limlar orasida davom etib ketmasin
            num, rest = s.split(".", 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(1)
            p.add_run(f"{num}.  ")
            _matn(p, rest.strip())
        else:
            p = doc.add_paragraph()
            if birinchi_paragraf:   # sarlavha ostidagi "xodimlar uchun · sana" qatori
                _matn(p, s, size=10, color=RGBColor(0x66, 0x66, 0x66))
                birinchi_paragraf = False
            else:
                _matn(p, s)
        i += 1
    doc.save(docx_yol)


def docx_dan_pdf(docx_yol, pdf_yol):
    try:
        import win32com.client
    except ImportError:
        print("  (pywin32 yo'q — PDF yaratilmadi)")
        return False
    w = win32com.client.DispatchEx("Word.Application")
    w.Visible = False
    try:
        d = w.Documents.Open(docx_yol)
        d.SaveAs2(pdf_yol, FileFormat=17)
        d.Close(False)
        return True
    except Exception as e:
        print(f"  PDF xatosi: {e}")
        return False
    finally:
        w.Quit()


def main(fayllar):
    if not fayllar:
        fayllar = sorted(f for f in os.listdir(PAPKA) if f.lower().endswith(".md"))
    os.makedirs(ONEDRIVE, exist_ok=True)
    for f in fayllar:
        md = os.path.join(PAPKA, os.path.basename(f))
        asos = os.path.splitext(md)[0]
        md_dan_docx(md, asos + ".docx")
        chiqdi = [asos + ".docx"]
        if docx_dan_pdf(asos + ".docx", asos + ".pdf"):
            chiqdi.append(asos + ".pdf")
        for c in [md] + chiqdi:
            shutil.copy2(c, ONEDRIVE)
        print(f"OK {os.path.basename(md)} -> {', '.join(os.path.basename(c) for c in chiqdi)}")
    print(f"OneDrive: {ONEDRIVE}")


if __name__ == "__main__":
    main(sys.argv[1:])
